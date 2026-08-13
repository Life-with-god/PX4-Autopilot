# -*- coding: utf-8 -*-
"""
生成 LOCP 使用 SOP Word 文档
解析 failure_detector_params.c 提取全部 LOCP_ 参数表，保证与代码一致
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'src', 'modules', 'commander', 'failure_detector',
                   'failure_detector_params.c')
OUT = os.path.join(HERE, '..', 'docs', 'zh', 'safety', 'LOCP_使用SOP.docx')

# 维度分组（按参数名前缀，按顺序）
GROUPS = [
    ('总开关', ['LOCP_EN']),
    ('ARD 姿态变化率', ['LOCP_ARD']),
    ('VRD 速度变化率', ['LOCP_VRD']),
    ('PRD 位置变化率', ['LOCP_PRD']),
    ('COD 电流异常', ['LOCP_COD']),
    ('MTO MAVLink超时', ['LOCP_MTO']),
    ('TRD 动力响应校验', ['LOCP_TRD']),
    ('HC 飞机健康门槛', ['LOCP_HC']),
    ('OBS Offboard异常', ['LOCP_OBS']),
    ('Crash 碰撞检测', ['LOCP_CRASH']),
]


def parse_params(path):
    """解析 params.c，返回 [(name, type, default, min, max, unit, desc), ...]"""
    text = open(path, encoding='utf-8').read()
    params = []
    pat = re.compile(
        r'/\*\*(.*?)\*/\s*PARAM_DEFINE_(FLOAT|INT32|BOOL)\(([A-Z0-9_]+),\s*([^)]+)\);',
        re.S)
    for m in pat.finditer(text):
        if not m.group(3).startswith('LOCP'):
            continue
        desc_lines = [l.strip().lstrip('*').strip() for l in m.group(1).split('\n')]
        desc_lines = [l for l in desc_lines if l and not l.startswith('@')]
        desc = ' '.join(desc_lines)
        # 提取元数据
        minv = re.search(r'@min\s+([0-9.\-]+)', m.group(1))
        maxv = re.search(r'@max\s+([0-9.\-]+)', m.group(1))
        unit = re.search(r'@unit\s+([A-Za-z/%²³]+)', m.group(1))
        params.append({
            'name': m.group(3),
            'type': m.group(2),
            'default': m.group(4).rstrip('f'),
            'min': minv.group(1) if minv else '',
            'max': maxv.group(1) if maxv else '',
            'unit': unit.group(1) if unit else '',
            'desc': desc,
        })
    return params


def set_font(run, name='微软雅黑', size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=16, bold=True, color=(0x1F, 0x4E, 0x79))
    p.space_before = Pt(0)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=13, bold=True, color=(0x1F, 0x4E, 0x79))
    return p


def add_h3(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=11.5, bold=True)
    return p


def add_para(doc, text, bold=False, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p


def add_bullet(doc, text, bold_head=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_head:
        r1 = p.add_run(bold_head)
        set_font(r1, bold=True)
    r = p.add_run(text)
    set_font(r)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    for line in text.strip().split('\n'):
        r = p.add_run(line)
        set_font(r, name='Consolas', size=9, color=(0x30, 0x30, 0x30))
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        p.add_run().add_break()
    return p


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        set_font(r, size=9.5, bold=True)
        shd = cell._element.get_or_add_tcPr().makeelement(qn('w:shd'), {
            qn('w:val'): 'clear', qn('w:fill'): 'DCE6F1'})
        cell._element.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v))
            set_font(r, size=9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def main():
    params = parse_params(SRC)
    print(f'解析 {len(params)} 个 LOCP 参数')

    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ============ 封面标题 ============
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('LOCP 失控保护系统\n标准作业流程（SOP）')
    set_font(r, size=22, bold=True, color=(0x1F, 0x4E, 0x79))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('适用：cuav_7-nano 定制固件（PX4 locp 分支）\n版本：v1.0    日期：2026-08-13')
    set_font(r, size=12)
    doc.add_paragraph()

    # ============ 一、系统概述 ============
    add_h1(doc, '一、系统概述')
    add_para(doc, 'LOCP（Loss-of-Control Protection，失控保护）通过 8 个独立维度实时检测飞行器失控，'
                  '并在检测到异常时自动执行停桨（Disarm）或降落（Land）保护动作。'
                  '所有检测默认禁用，需在 QGC 中手动开启。')
    add_h2(doc, '1.1 检测维度')
    add_table(doc, ['维度', '英文全称', '含义', '数据源'], [
        ['ARD', 'Attitude Rate Detection', '姿态变化率异常（角加速度尖峰 / 持续高角速率）', 'IMU 角速度'],
        ['VRD', 'Velocity Rate Detection', '速度变化率异常（水平加速度 / 自由落体 / jerk / 水平速度持续异常）', 'EKF local_position'],
        ['PRD', 'Position Rate Detection', '位置变化率异常（急降 / 高度振荡 / 水平漂移）', 'EKF local_position'],
        ['COD', 'Current Overdraw Detection', '电流异常（总电流突增 / dI/dt 尖峰）', 'battery_status'],
        ['MTO', 'MAVLink Timeout', 'MAVLink 消息超时（心跳 / 指令 / 速率骤降）', 'MAVLink 链路'],
        ['OBS', 'Offboard Setpoint Sanity', 'Offboard Setpoint 异常（数值跳变 / NaN 注入）', 'trajectory_setpoint'],
        ['TRD', 'Thrust Response Detection', '动力响应校验（高油门但无运动响应，卡网/动力丢失）', '油门 + EKF 运动状态'],
        ['Crash', 'Crash/Impact Detection', '碰撞/撞击检测（IMU 加速度范数尖峰，无迟滞）', 'IMU 加速度'],
    ], widths=[1.7, 4.6, 6.6, 3.1])
    add_h2(doc, '1.2 保护动作分配（固定动作，无等级参数）')
    add_table(doc, ['维度', '飞机自身状态', '动作'], [
        ['ARD / VRD / PRD / COD / Crash', '已失控（控制回路不可信）', '停桨 Disarm'],
        ['TRD', '动力无响应（卡网等）', '分级：低高度先降落，超时/失联/高高度转停桨'],
        ['MTO / OBS', '仅外部输入异常（飞机可受控）', '健康检查通过 → 降落 Land'],
        ['MTO / OBS', '外部异常 + 飞机不健康', '停桨 Disarm'],
    ], widths=[5.5, 6.0, 4.2])
    p = add_para(doc, '核心原则：降落的前提是飞机自身受控；飞机自身失控一律停桨。'
                      'MTO/OBS 触发时由 checkVehicleHealthy() 实时检查角速率/水平速度/垂直速度/数据有效性，'
                      '任何一项超门槛（LOCP_HC_*）即判定不健康 → 停桨。', bold=True)
    p.paragraph_format.space_before = Pt(6)

    # ============ 二、前置条件 ============
    add_h1(doc, '二、前置条件')
    add_table(doc, ['项目', '要求'], [
        ['硬件', 'CUAV 7-Nano 飞控 + USB Type-C 数据线'],
        ['固件', 'locp 分支编译的 cuav_7-nano_default.px4（当前 firmware/ 目录下）'],
        ['软件', 'QGroundControl（参数配置/烧录/日志）、Python 3 + pyserial（命令行烧录）'],
        ['环境', 'Windows 10/11 或 WSL（Ubuntu）'],
    ], widths=[3.0, 12.7])

    # ============ 三、固件烧录 ============
    add_h1(doc, '三、固件烧录 SOP')
    add_h2(doc, '3.1 Bootloader 概念（一句话）')
    add_para(doc, 'Bootloader 是存放在 Flash 前 128KB 的独立小程序，上电时负责引导固件，'
                  '刷机时负责擦写应用区。刷固件不会覆盖它，固件刷坏也能重新再刷。'
                  '日常刷机通过 PX4 BL（设备显示 "PX4 BL CUAV 7 Nano"）完成。')
    add_h2(doc, '3.2 进入 Bootloader 模式（三选一）')
    add_table(doc, ['方式', '操作', '适用场景'], [
        ['① QGC 自动', 'QGC 齿轮图标 → Firmware → 选择自定义固件文件 → QGC 自动重启飞控进入 BL 并烧录', '最常用（推荐）'],
        ['② 安全开关', '按住安全开关不放再上电，蜂鸣器长响后松开，飞控进入 BL', '无 QGC 时'],
        ['③ NSH 命令', 'MAVLink Shell 执行 reboot -b', '已连接 Shell 时'],
    ], widths=[2.6, 9.5, 3.6])
    add_h2(doc, '3.3 命令行烧录（Windows）')
    add_para(doc, '重要：烧录前必须关闭 QGroundControl（否则 COM 口被占用，报"拒绝访问"）。', bold=True)
    add_code(doc, '''# ① 确认进入 BL（设备管理器出现 "PX4 BL CUAV 7 Nano"，记下 COM 号）
# ② 关闭 QGroundControl
# ③ 执行烧录（COM3 替换为实际端口号）：
python Tools/px_uploader.py --port COM3 firmware/cuav_7-nano_default.px4

# 成功标志：Erase 100% / Program 100% / Verify 100% / Rebooting''')
    add_para(doc, '注：WSL 中端口为 /dev/ttyACM0 形式，命令相同。')
    add_h2(doc, '3.4 烧录后验证')
    add_code(doc, '''# MAVLink Shell 中：
ver                     # 版本信息包含 locp
param show LOCP_EN      # 能查到参数即说明 LOCP 已编译进固件''')
    add_h2(doc, '3.5 BOOT 键 DFU 模式（仅救砖，不用于日常刷机）')
    add_para(doc, '按住板上 BOOT 键上电进入的是芯片 ROM 内置 DFU 模式（设备显示 "STM32 BOOTLOADER"），'
                  '此时 px_uploader 无法使用，需 STM32CubeProgrammer 烧录 bootloader 本体（地址 0x08000000）。'
                  '仅当 PX4 BL 损坏时使用。')

    # ============ 四、参数配置 ============
    add_h1(doc, '四、参数配置 SOP')
    add_h2(doc, '4.1 两级配置方案')
    add_table(doc, ['方案', '启用维度', '适合阶段'], [
        ['保守配置（推荐首飞）', 'LOCP_EN + CRASH + TRD', '首飞验证期：先验证无误报，再逐项放开'],
        ['完整配置', '全部 8 维度', '验证通过后的日常作业'],
    ], widths=[5.0, 5.5, 5.2])
    add_h2(doc, '4.2 配置方法')
    add_para(doc, 'QGC → 齿轮图标 → Parameters → 搜索 LOCP → 修改参数 → Tools → 全部保存到闪存。'
                  '也可在 MAVLink Shell 中粘贴 4.3 脚本后执行 param save。')
    add_h2(doc, '4.3 一键参数脚本（完整配置）')
    add_para(doc, '注意：EN 开关需连同子开关一起开启；HC 健康门槛参数无需开启开关，始终生效。', bold=True)
    add_code(doc, '''# === 总开关 + 碰撞 + 动力响应（安全底线）===
param set LOCP_EN 1
param set LOCP_CRASH_EN 1
param set LOCP_TRD_EN 1

# === ARD 姿态变化率 ===
param set LOCP_ARD_EN 1 ; param set LOCP_ARD_ACC_EN 1 ; param set LOCP_ARD_RATE_EN 1
param set LOCP_ARD_R_MAX 100 ; param set LOCP_ARD_P_MAX 120 ; param set LOCP_ARD_Y_MAX 60
param set LOCP_ARD_RSP 2 ; param set LOCP_ARD_PSP 2 ; param set LOCP_ARD_YSP 5
param set LOCP_ARD_DUR 0.2 ; param set LOCP_ARD_T 0.1

# === VRD 速度变化率 ===
param set LOCP_VRD_EN 1 ; param set LOCP_VRD_AH_EN 1 ; param set LOCP_VRD_FF_EN 1
param set LOCP_VRD_JK_EN 1 ; param set LOCP_VRD_HS_EN 1
param set LOCP_VRD_AH_MAX 55 ; param set LOCP_VRD_AD_MAX 6 ; param set LOCP_VRD_VZD_MAX 3
param set LOCP_VRD_JERK 100 ; param set LOCP_VRD_T 0.3
param set LOCP_VRD_HS_MAX 15 ; param set LOCP_VRD_HS_DUR 3.0

# === PRD 位置变化率 ===
param set LOCP_PRD_EN 1 ; param set LOCP_PRD_DES_EN 1 ; param set LOCP_PRD_OSC_EN 1
param set LOCP_PRD_DRF_EN 1
param set LOCP_PRD_VZ_MAX 5 ; param set LOCP_PRD_ADROP 3
param set LOCP_PRD_ASTD 2 ; param set LOCP_PRD_HSPD 5 ; param set LOCP_PRD_T 0.5

# === COD 电流异常 ===
param set LOCP_COD_EN 1 ; param set LOCP_COD_SRG_EN 1 ; param set LOCP_COD_SPK_EN 1
param set LOCP_COD_MAX_I 100 ; param set LOCP_COD_DELTA_I 10
param set LOCP_COD_DI_DT 300 ; param set LOCP_COD_T 0.5 ; param set LOCP_COD_ARM_DLY 2.0

# === MTO MAVLink 超时 ===
param set LOCP_MTO_EN 1 ; param set LOCP_MTO_HB_EN 1 ; param set LOCP_MTO_CMD_EN 1
param set LOCP_MTO_HB_T 1.5 ; param set LOCP_MTO_CMD_T 2.0 ; param set LOCP_MTO_RATE 3

# === OBS Offboard 异常 ===
param set LOCP_OBS_EN 1 ; param set LOCP_OBS_JMP_EN 1 ; param set LOCP_OBS_NAN_EN 1
param set LOCP_OBS_J_POS 10 ; param set LOCP_OBS_J_VEL 5 ; param set LOCP_OBS_J_YAW 1.57
param set LOCP_OBS_T 0.3

param save''')
    add_h2(doc, '4.4 保守配置脚本（首飞推荐）')
    add_code(doc, '''param set LOCP_EN 1
param set LOCP_CRASH_EN 1
param set LOCP_TRD_EN 1
param save''')

    # ============ 五、日常使用 ============
    add_h1(doc, '五、日常使用 SOP')
    add_h2(doc, '5.1 飞行前检查清单')
    for txt in ['总开关 LOCP_EN=1 已确认（避免"配了参数忘了开总开关"）',
                'QGC Parameters 页确认各项 EN 开关与阈值符合预期',
                'MAVLink Inspector 观察 failsafe_flags：所有 locp_*_triggered=false、crash_detected=false',
                '螺旋桨、电机、机架无损伤（LOCP 停桨是最后防线，不替代日常维护）']:
        add_bullet(doc, txt, '□ ')
    add_h2(doc, '5.2 飞行中监控')
    add_bullet(doc, 'QGC → MAVLink Inspector → 搜索 failsafe_flags，实时观察 locp_ard_triggered、'
                    'locp_vrd_triggered、locp_trd_triggered、crash_detected 等字段')
    add_bullet(doc, 'locp_vehicle_healthy 字段应始终为 true（飞机自身健康检查）')
    add_h2(doc, '5.3 飞行后日志分析')
    for txt in ['导出 SD 卡 .ulg 日志 → https://review.px4.io/ 或 PlotJuggler',
                '搜索 failsafe_flags 话题，核对各 locp_*_triggered 时间点与飞行录像对照',
                '重点确认：正常机动段无触发（无误报）、已知异常段有触发（无漏报）']:
        add_bullet(doc, txt)
    add_h2(doc, '5.4 触发后处置')
    add_table(doc, ['触发维度', '动作已发生', '飞行后处置'], [
        ['ARD/VRD/PRD/COD/Crash', '停桨（空中坠机/地面停转）', '检查桨叶/机架/电机，排除故障后重新解锁'],
        ['TRD 低高度', '先尝试降落，超时/失联转停桨', '检查是否卡网/桨损坏/动力丢失'],
        ['MTO/OBS 健康', '降落', '检查数传链路 / 机载计算机程序'],
        ['MTO/OBS 不健康', '停桨', '同时存在飞机自身失控，按失控流程全面检查'],
    ], widths=[3.8, 5.2, 6.7])

    # ============ 六、调参与常见问题 ============
    add_h1(doc, '六、调参与常见问题')
    add_h2(doc, '6.1 误触发处理')
    add_table(doc, ['现象', '调整'], [
        ['正常飞行中 ARD 误触发', '增大 LOCP_ARD_R_MAX/P_MAX 或 LOCP_ARD_T'],
        ['正常着陆触发 Crash', '增大 LOCP_CRASH_THR（已标定 70，正常触地最高 61.8）'],
        ['解锁瞬间 COD 误触发', '增大 LOCP_COD_ARM_DLY'],
        ['数传偶发丢包触发 MTO', '增大 LOCP_MTO_HB_T'],
        ['正常快速机动触发 VRD_HS', '增大 LOCP_VRD_HS_MAX 或 LOCP_VRD_HS_DUR'],
    ], widths=[6.5, 9.2])
    add_h2(doc, '6.2 标定依据（2026-08-13，6-8 月 136 份纯净日志）')
    add_para(doc, '全部阈值按本机日志标定：正常样本裕度 1.3~11 倍；'
                  '6-8 月 17 份失控日志（EGO 故障 13 / 卡网 2 / 翻倒 1 / 碰撞 1）全部被现有维度覆盖。'
                  '阈值已按"正常与失控取中点"原则设定，不建议无依据修改。')
    add_h2(doc, '6.3 已知注意点')
    for txt in ['EGO 视觉故障：水平速度发散到 47~107 m/s，依赖 LOCP_VRD_HS_EN 通道捕获，请保持其开启',
                '地面假速度：地面 arm 调试时 EGO 可能给 18~22 m/s 假速度，此时飞机在地面，停桨无害',
                '卡网场景水平速度低（<4 m/s），由 TRD 动力响应校验兜底，请保持 LOCP_TRD_EN 开启',
                '1-2 月地面翻倒日志属失控样本，勿用其数据重新标定阈值']:
        add_bullet(doc, txt)

    # ============ 七、安全注意事项 ============
    add_h1(doc, '七、安全注意事项')
    for txt in ['所有 EN 开关默认禁用（=0），必须手动开启后才生效',
                '停桨（Disarm）动作不可逆：空中停桨即坠机，请确保阈值经过实飞验证',
                '首飞务必采用保守配置（仅 CRASH + TRD），确认无误报后逐项放开',
                'MTO/OBS 的降落动作以飞机自身健康检查通过为前提（LOCP_HC_* 门槛）',
                '刷固件不会清除参数，但换板/恢复出厂会重置，配置后请导出参数备份']:
        add_bullet(doc, txt)

    # ============ 八、参数速查表 ============
    add_h1(doc, '八、参数速查表（70 个，与固件代码一致）')
    used = set()
    for gname, prefixes in GROUPS:
        items = [p for p in params if p['name'] not in used
                 and any(p['name'].startswith(prefix) for prefix in prefixes)]
        used.update(p['name'] for p in items)
        if not items:
            continue
        add_h3(doc, f'8.{GROUPS.index((gname, prefixes)) + 1} {gname}（{len(items)} 个参数）')
        add_table(doc, ['参数', '默认值', '范围', '单位', '说明'], [
            [p['name'], p['default'],
             f"{p['min']}~{p['max']}" if p['min'] and p['max'] else '—',
             p['unit'] or '—', p['desc']]
            for p in items], widths=[3.6, 1.5, 1.8, 1.4, 7.4])

    doc.save(OUT)
    print(f'SOP Word 已生成: {os.path.abspath(OUT)}')


if __name__ == '__main__':
    main()
